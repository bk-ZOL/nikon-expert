"""src/ts_table.py —— Troubleshooting 积累表（2026-10-15）。

两层，格式照《苏州星钥S207D现场调试记录.xlsx》的「报警处理记录」：
  案例 ts_cases：一条报警一行 —— 编号 | 登记日期 | 机型 | 客户/现场 | Unit | 报警/问题描述 | 故障代码 |
                 排查过程（试了什么、结果怎样，没用的也写）| 解决措施（只写真正解决问题的那一步）|
                 Root cause | 根因分类 | 更换备件 | 状态(Open/Pending/Close) | 处理人 | 关闭日期 | 后续计划/注意事项
                 来源：故障履历 / 工作日志 / 客户交付物里用模型抽取；或人工整理好的表直接导入（算已确认）。
  条目 ts_entries：按 机型 → Unit → 故障代码（无代码按现象）归并 —— 发生次数和日期、根因排序、有效解决措施排序、
                 走过的弯路（排查时无效的尝试）、更换过的备件、原厂对策。智能体查的是这一层。

发生次数：同一份资料里相隔 ≤3 天的记录算一次（日志每天往下抄）；隔得久的算新的一次；不同资料（不同现场）分开算。
审核：抽取出的案例和条目先「待确认」，管理员在网页确认后才给其他账号看、才被智能体引用；
      人工改过的条目不被自动更新覆盖，有新案例时标 has_new 提醒复核。
权限：需要「故障履历」类别 + 对应机型；客户来源的案例只给有该客户权限的人看。
数据在 data/ts_table.db（独立库，不和全文检索抢锁）。
"""
import hashlib
import json
import os
import re
import time

try:
    import pysqlite3 as sqlite3
except ImportError:
    import sqlite3

DB = os.getenv("TS_TABLE_DB", "./data/ts_table.db")

# Unit（照参考表的下拉选项，补了 AF / AIS / MAC / 干涉计）。抽取时只能从这里选。
UNITS = ["WS", "RS", "RB", "WL", "RL", "LC", "AVIS", "ALG", "AF", "AIS", "BSA", "Illumination", "BMU", "Laser",
         "Chamber", "Load chamber", "Inline", "MAC", "干涉计", "主机/PC", "电源/厂务", "接线", "其他"]
RC_CATEGORIES = ["接头/线路", "硬件损坏", "机械/调整", "参数/设定", "厂务/二次配", "电源/电压", "通讯", "软件", "待查"]
STATUSES = ["Open", "Pending", "Close"]

SOURCE_CATEGORIES = ("fault_history", "worklog", "customer_deliverable")
WINDOW_CHARS = 2600          # 每次送给模型的原文长度
MIN_CONTENT_CHARS = 40       # 去掉模板字样后不足这么多字的段落跳过（空模板）
OCC_GAP_DAYS = 3             # 相隔这么多天以内的记录算同一次发生
_TEMPLATE_WORDS = re.compile(r"故障日期|机台|Error Code|故障现象|根因分析|解决方案|处理人员|[:：()（）\s]")

CASE_COLS = ["case_no", "reg_date", "machine", "site", "unit", "description", "error_code", "process",
             "solution", "root_cause", "rc_category", "parts", "status", "handler", "close_date", "follow_up"]
CASE_LABELS = ["编号", "登记日期", "机型", "客户/现场", "Unit", "报警/问题描述", "故障代码", "排查过程",
               "解决措施", "Root cause", "根因分类", "更换备件", "状态", "处理人", "关闭日期", "后续计划/注意事项"]


def _conn():
    c = sqlite3.connect(DB)
    c.row_factory = sqlite3.Row
    c.executescript(f"""
    CREATE TABLE IF NOT EXISTS ts_cases (
        id INTEGER PRIMARY KEY AUTOINCREMENT, entry_id INTEGER,
        {", ".join(f"{k} TEXT DEFAULT ''" for k in CASE_COLS)},
        symptom_key TEXT DEFAULT '', customer TEXT DEFAULT 'none', dates TEXT DEFAULT '[]',
        src_doc TEXT, src_section TEXT, dedup_key TEXT, review TEXT DEFAULT 'pending', created_at TEXT);
    CREATE UNIQUE INDEX IF NOT EXISTS ux_case_dedup ON ts_cases(src_doc, dedup_key);
    CREATE TABLE IF NOT EXISTS ts_entries (
        id INTEGER PRIMARY KEY AUTOINCREMENT, entry_key TEXT UNIQUE,
        machine TEXT, unit TEXT, error_code TEXT, title TEXT,
        root_causes TEXT DEFAULT '[]', solutions TEXT DEFAULT '[]', detours TEXT DEFAULT '[]',
        parts TEXT DEFAULT '[]', follow_up TEXT DEFAULT '', oem_refs TEXT DEFAULT '[]',
        occ_count INTEGER DEFAULT 0, occ_dates TEXT DEFAULT '[]', first_date TEXT, last_date TEXT,
        open_count INTEGER DEFAULT 0, customers TEXT DEFAULT '',
        status TEXT DEFAULT 'pending', edited INTEGER DEFAULT 0, has_new INTEGER DEFAULT 0,
        notes TEXT DEFAULT '', edited_by TEXT, updated_at TEXT);
    CREATE TABLE IF NOT EXISTS ts_windows (hash TEXT PRIMARY KEY, src_doc TEXT, n_cases INTEGER, done_at TEXT);
    """)
    cols = {r[1] for r in c.execute("PRAGMA table_info(ts_entries)")}
    for col, ddl in (("origin", "TEXT DEFAULT 'field'"), ("oem_detail", "TEXT DEFAULT '[]'"),
                     ("src", "TEXT DEFAULT ''"), ("match_key", "TEXT DEFAULT ''"), ("exp_mix", "TEXT DEFAULT ''")):
        if col not in cols:
            c.execute(f"ALTER TABLE ts_entries ADD COLUMN {col} {ddl}")
    if "exp_type" not in {r[1] for r in c.execute("PRAGMA table_info(ts_cases)")}:
        c.execute("ALTER TABLE ts_cases ADD COLUMN exp_type TEXT DEFAULT ''")
        for r in c.execute("SELECT id, src_doc FROM ts_cases").fetchall():
            c.execute("UPDATE ts_cases SET exp_type=? WHERE id=?", (exp_type_of(r["src_doc"]), r["id"]))
        c.commit()
    c.execute("CREATE INDEX IF NOT EXISTS ix_entry_match ON ts_entries(machine, match_key)")
    return c


# 经验类型：厂内翻新（翻新 / 改造 / 厂内调试）、现场（客户现场装机、故障处理）；原厂推荐来自 Nikon Guide（origin=oem）
EXP_REFURB, EXP_FIELD = "翻新", "现场"
_REFURB_WORDS = [w for w in os.getenv("TS_REFURB_KEYWORDS", "翻新,厂内,改造,refurb").split(",") if w]


def exp_type_of(src_doc: str) -> str:
    return EXP_REFURB if any(w.lower() in (src_doc or "").lower() for w in _REFURB_WORDS) else EXP_FIELD


def _now():
    return time.strftime("%Y-%m-%d %H:%M:%S")


def norm_code(code: str) -> str:
    code = (code or "").strip().replace("：", ":")
    return "" if code.upper() in ("", "-", "—", "无", "NONE", "N/A") else code


def _norm_text(s: str) -> str:
    return re.sub(r"[\s，,。．.；;：:、()（）\-_/]+", "", (s or "").lower())


def _code_key(code: str) -> str:
    """多个代码时取第一个、去空格大写，作归并用。"""
    first = re.split(r"[；;／/]", code or "")[0]
    return re.sub(r"\s+", "", first).upper()[:60]


def match_key(code: str) -> str:
    """跨来源对错误码：STG-4A-1B / 4A-1B / #4AH-1B → 4A-1B；8 位码原样；没有码返回空。"""
    k = _code_key(code)
    k = re.sub(r"^(STG|ST|WS|RS|WL|RL|AF|LC|MAC)[-_:]", "", k)
    k = k.lstrip("#")
    m = re.match(r"^([0-9A-F]{2})H?[-_]?([0-9A-FX]{2})H?$", k)
    if m:
        return f"{m.group(1)}-{m.group(2)}"
    return k


def _parse_json(text: str, kind="list"):
    pat = r"\[.*\]" if kind == "list" else r"\{.*\}"
    m = re.search(pat, text or "", re.S)
    if not m:
        return [] if kind == "list" else {}
    try:
        return json.loads(m.group(0))
    except Exception:
        return [] if kind == "list" else {}


def _llm(prompt: str) -> str:
    """直接调 OpenAI 兼容接口（deepseek）。默认思考档会把 max_tokens 全用在思考上、正文为空，
    所以思考档设 low（TS_REASONING，off=关闭思考），输出上限放宽（TS_MAX_TOKENS）。
    TS_LLM=local：改用 Mac 上的本地模型（src/local_llm.py，资料不出自己的设备）；Mac 不可用就报错，不退回云端。"""
    if os.getenv("TS_LLM", "").lower() == "local":
        from src import local_llm
        ok, why = local_llm.available(require_idle=False)
        if not ok:
            raise RuntimeError(f"本地模型不可用：{why}")
        return local_llm.chat(prompt)
    from llama_index.core import Settings
    from openai import OpenAI
    llm = Settings.llm
    cli = OpenAI(api_key=getattr(llm, "api_key", None), base_url=getattr(llm, "api_base", None), timeout=300)
    kw = {}
    effort = os.getenv("TS_REASONING", "low")
    if effort == "off":
        kw["extra_body"] = {"thinking": {"type": "disabled"}}
    elif effort:
        kw["reasoning_effort"] = effort
    r = cli.chat.completions.create(model=llm.model, messages=[{"role": "user", "content": prompt}],
                                    max_tokens=int(os.getenv("TS_MAX_TOKENS", "16000")), **kw)
    return r.choices[0].message.content or ""


# ── 1. 原文窗口 ───────────────────────────────────────────────────────
def source_windows(only_doc: str = None):
    """[(src_doc, machine, customer, text, section)]：履历/日志/客户交付物按段拼成 ~2600 字的窗口。"""
    fts = sqlite3.connect(os.getenv("FTS_DB_PATH", "./data/fts.db"))
    docs = fts.execute(f"""
        SELECT filename, MAX(models), MAX(customer) FROM doc_acl
        WHERE category IN ({",".join("?" * len(SOURCE_CATEGORIES))}) GROUP BY filename""",
                       SOURCE_CATEGORIES).fetchall()
    for name, models, cust in docs:
        if only_doc and only_doc not in name:
            continue
        machine = next((m for m in (models or "").split(",") if m and m != "ALL"), "S207D")
        rows = fts.execute("SELECT section_title, text FROM doc_fts WHERE doc_name=? ORDER BY rowid", (name,)).fetchall()
        buf, secs = "", []
        for sec, text in rows:
            text = (text or "").strip()
            if len(_TEMPLATE_WORDS.sub("", text)) < MIN_CONTENT_CHARS:
                continue
            if buf and len(buf) + len(text) > WINDOW_CHARS:
                yield name, machine, cust or "none", buf, " / ".join(secs[:3])
                buf, secs = "", []
            buf += f"\n【{sec}】\n{text}" if sec else f"\n{text}"
            if sec:
                secs.append(sec.strip("# ")[:40])
        if buf:
            yield name, machine, cust or "none", buf, " / ".join(secs[:3])
    fts.close()


# ── 2. 抽取案例 ───────────────────────────────────────────────────────
EXTRACT_PROMPT = """你是 Nikon 光刻机（NSR 系列）的资深维修专家。下面是现场的故障履历 / 工作日志原文（机型：{machine}）。
请整理成「报警处理记录」：一条报警 / 故障一行。重点是**故障根因（Root cause）**和**真正解决问题的措施**。

规则：
1. 每个独立的报警/故障单独一条（一段里常同时有好几个故障，要拆开）；只有计划、进度、施工协调、没有故障的内容不要抽。
2. 同一段里重复出现的同一个故障只抽一次。
3. 排查过程（process）：试了什么、结果怎样，**没用的尝试也写**（下次能少走弯路）。
4. 解决措施（solution）：**只写最后真正解决问题的那一步**，具体到部件、参数、数值（如「RB 压力 0.278→0.29 后复位通过」）。
   「检查 / 确认 / 观察」这类排查步骤，除非它本身就解决了问题，否则不算解决措施。还没解决就留空。
5. Root cause：查明的原因（不是现象：「报警 STG-B8-E0」是现象，「RB 供气压力偏低」才是根因）；只是推测就在前面加「推定：」；没查出来写「待查」。
6. 故障代码照原文抄，多个用「；」隔开；没有写空字符串。
   报警界面原文「STG - FATAL - E Detail: 60730410 FFF0FE17 …」要换算：第 1、2 串各取最后两位 → 写 STG-10-17。
7. 状态：Close = 根因确认并解决；Pending = 临时措施已上 / 在等备件或条件；Open = 还没头绪。

Unit 只能从这里选一个：{units}
根因分类只能从这里选一个：{rcs}

只输出 JSON 数组（没有就输出 []），不要注释，每个元素：
{{"date": "YYYY-MM-DD（看不出就空）", "unit": "Unit", "description": "报警/问题描述，一句话",
  "symptom_key": "现象规范化短语（≤12字，用于把相同故障归到一起，如「FFU 温控停止」「CT 轴无法升起」）",
  "error_code": "故障代码", "process": "排查过程", "solution": "解决措施", "root_cause": "Root cause",
  "rc_category": "根因分类", "parts": "更换的备件（没有空）", "status": "Open|Pending|Close",
  "close_date": "YYYY-MM-DD（Close 时能看出就填）", "follow_up": "后续计划/注意事项（没有空）"}}

原文：
{text}
"""


def extract_cases(machine: str, text: str) -> list:
    out = _parse_json(_llm(EXTRACT_PROMPT.format(
        machine=machine, units="、".join(UNITS), rcs="、".join(RC_CATEGORIES), text=text[:WINDOW_CHARS + 400])))
    cases = []
    for c in out if isinstance(out, list) else []:
        if not isinstance(c, dict) or not (c.get("description") or c.get("error_code")):
            continue
        c = {k: str(v or "").strip() for k, v in c.items()}
        c["unit"] = c.get("unit") if c.get("unit") in UNITS else "其他"
        c["rc_category"] = c.get("rc_category") if c.get("rc_category") in RC_CATEGORIES else "待查"
        c["status"] = c.get("status") if c.get("status") in STATUSES else "Open"
        c["error_code"] = norm_code(c.get("error_code"))
        c["reg_date"] = c.pop("date", "")
        cases.append(c)
    return cases


# ── 3. 写入案例、归并条目 ─────────────────────────────────────────────
def entry_key(machine, unit, code, symptom_key) -> str:
    ck = match_key(code)        # 4C-04 / STG-4C-04 / ST-4C-04 归到同一个条目
    return f"{machine}|{unit}|" + (f"code:{ck}" if ck else f"sym:{_norm_text(symptom_key)[:30]}")


def remerge(log=print) -> int:
    """按新的 entry_key 重新归并现场条目（同码不同写法被拆开的合起来），合并后重新整理。返回合并掉的条目数。"""
    c = _conn()
    rows = c.execute("SELECT * FROM ts_entries WHERE COALESCE(origin,'field')='field' ORDER BY id").fetchall()
    groups = {}
    for e in rows:
        cs = c.execute("SELECT symptom_key, description FROM ts_cases WHERE entry_id=? LIMIT 1", (e["id"],)).fetchone()
        sym = (cs["symptom_key"] or cs["description"]) if cs else e["title"]
        groups.setdefault(entry_key(e["machine"], e["unit"], e["error_code"], sym), []).append(e)
    merged, touched = 0, set()
    for k, es in groups.items():               # 第一遍：同组的合并到最早的那个条目（人工改过的优先保留）
        es.sort(key=lambda e: (-(e["edited"] or 0), e["id"]))
        tgt = es[0]["id"]
        for e in es[1:]:
            c.execute("UPDATE ts_cases SET entry_id=? WHERE entry_id=?", (tgt, e["id"]))
            c.execute("DELETE FROM ts_entries WHERE id=?", (e["id"],))
            touched.add(tgt)
            merged += 1
    for k, es in groups.items():               # 第二遍：改成新键（重复的已删，不会撞唯一约束）
        c.execute("UPDATE ts_entries SET entry_key=? WHERE id=?", (k, es[0]["id"]))
    c.commit()
    refresh_entries(c, touched, log)
    backfill_match_keys(c)
    c.close()
    log(f"重新归并：合并掉 {merged} 个重复条目，重新整理 {len(touched)} 个")
    return merged


def add_case(c, machine: str, case: dict, src_doc: str, section: str, customer: str, review="pending") -> tuple:
    """写入一条案例并挂到条目上。返回 (entry_id, 是否有变化)。
    同一份资料里同一个故障（同 Unit + 代码/现象 + 解决措施）只存一条，再出现只把日期记上。"""
    sym = case.get("symptom_key") or case.get("description", "")
    dedup = hashlib.md5((case.get("unit", "") + _code_key(case.get("error_code")) + _norm_text(sym)[:40]
                         + _norm_text(case.get("solution"))[:80]).encode()).hexdigest()
    key = entry_key(machine, case.get("unit", "其他"), case.get("error_code"), sym)
    row = c.execute("SELECT id FROM ts_entries WHERE entry_key=?", (key,)).fetchone()
    if row:
        eid = row["id"]
    else:
        eid = c.execute("""INSERT INTO ts_entries(entry_key, machine, unit, error_code, title, updated_at, match_key)
                           VALUES (?,?,?,?,?,?,?)""",
                        (key, machine, case.get("unit", "其他"), case.get("error_code", ""),
                         case.get("description") or sym, _now(), match_key(case.get("error_code", "")))).lastrowid
    date = (case.get("reg_date") or "")[:10]
    vals = {k: case.get(k, "") for k in CASE_COLS}
    vals.update(machine=machine, reg_date=date)
    try:
        c.execute(f"""INSERT INTO ts_cases(entry_id, {", ".join(CASE_COLS)}, symptom_key, customer, dates,
                      src_doc, src_section, dedup_key, review, created_at, exp_type)
                      VALUES ({",".join("?" * (len(CASE_COLS) + 10))})""",
                  [eid] + [vals[k] for k in CASE_COLS] + [sym, customer, json.dumps([date] if date else []),
                                                          src_doc, section, dedup, review, _now(),
                                                          exp_type_of(src_doc)])
        return eid, True
    except sqlite3.IntegrityError:
        if date:
            r = c.execute("SELECT id, dates FROM ts_cases WHERE src_doc=? AND dedup_key=?", (src_doc, dedup)).fetchone()
            ds = sorted(set(json.loads(r["dates"] or "[]")) | {date})
            c.execute("UPDATE ts_cases SET dates=? WHERE id=?", (json.dumps(ds), r["id"]))
            return eid, True
        return eid, False


def occurrences(cases) -> list:
    """发生次数：同一现场相隔 ≤3 天的日期合成一次；没日期的每条算一次。返回每次的起始日期。
    「同一现场」：客户资料按客户（人工整理表和日报都来自同一现场，同一次故障不重复算）；内部资料按资料名。"""
    from datetime import date
    out, by_doc = [], {}
    for x in cases:
        ds = json.loads(x["dates"] or "[]") or ([x["reg_date"]] if x["reg_date"] else [])
        if not ds:
            out.append("")
            continue
        site = x["customer"] if x["customer"] not in ("", "none") else x["src_doc"]
        by_doc.setdefault(site, set()).update(ds)
    for ds in by_doc.values():
        prev = None
        for d in sorted(ds):
            try:
                cur = date.fromisoformat(d[:10])
            except ValueError:
                out.append(d)
                continue
            if prev is None or (cur - prev).days > OCC_GAP_DAYS:
                out.append(d[:10])
            prev = cur
    return sorted(out, key=lambda d: d or "9999")


CONSOLIDATE_PROMPT = """下面是 Nikon {machine}「{unit}」同一类故障{code}的 {n} 条处理记录，每条前面标了经验来源：
【现场】= 客户现场装机 / 故障处理；【翻新】= 厂内翻新 / 改造 / 调试。请整理成 Troubleshooting 汇总的一行：
- root_causes：**一个报警可能有多个根因，全部列出**（不是现象），意思相同的合并；按出现次数从多到少；只是推测的标「推定」；
  field / refurb = 这个根因在现场 / 翻新记录里各出现几次；category 只能从这里选：{rcs}；
- solutions：真正解决过问题的措施，意思相同的合并；field / refurb = 在现场 / 翻新时各解决过几次；写明对应哪个根因（for_cause 填根因原文）；
- detours：排查时试过但没用的动作（下次可以少走的弯路），合并同义的；
- process：建议的处理顺序，先做解决次数多、操作简单的。
只输出一个 JSON 对象，不要注释：
{{"title": "报警/问题描述，一句话",
  "root_causes": [{{"cause": "根因", "category": "根因分类", "field": 1, "refurb": 0}}],
  "solutions": [{{"action": "解决措施（具体）", "for_cause": "对应的根因", "field": 1, "refurb": 0}}],
  "detours": ["试过无效的动作"],
  "process": ["第一步…", "第二步…"]}}
记录：
{cases}
"""


def consolidate(c, eid: int):
    e = c.execute("SELECT * FROM ts_entries WHERE id=?", (eid,)).fetchone()
    cases = c.execute("SELECT * FROM ts_cases WHERE entry_id=? ORDER BY reg_date, id", (eid,)).fetchall()
    if not cases:
        return
    occ = occurrences(cases)
    dates = sorted(d for x in cases for d in (json.loads(x["dates"] or "[]") or ([x["reg_date"]] if x["reg_date"] else [])))
    custs = sorted({x["customer"] for x in cases if x["customer"] not in ("", "none")})
    parts = sorted({p.strip() for x in cases for p in re.split(r"[；;、,，]", x["parts"] or "") if p.strip()})
    meta = dict(occ_count=len(occ), occ_dates=json.dumps(occ, ensure_ascii=False),
                first_date=dates[0] if dates else "", last_date=dates[-1] if dates else "",
                open_count=sum(1 for x in cases if x["status"] != "Close"),
                customers=",".join(custs), parts=json.dumps(parts, ensure_ascii=False), updated_at=_now())
    if all(x["review"] == "confirmed" for x in cases) and e["status"] == "pending" and not e["edited"]:
        meta["status"] = "confirmed"        # 全部来自人工整理好的表 → 直接算确认
    if e["edited"]:
        meta["has_new"] = 1
        c.execute(f"UPDATE ts_entries SET {', '.join(k + '=?' for k in meta)} WHERE id=?", list(meta.values()) + [eid])
        return
    lines = []
    for i, x in enumerate(cases[:25], 1):
        lines.append(f"{i}. 【{x['exp_type'] or exp_type_of(x['src_doc'])}】[{x['reg_date'] or '日期不明'}｜{x['status']}] 描述：{x['description']}｜代码：{x['error_code'] or '—'}"
                     f"｜排查过程：{x['process'] or '—'}｜解决措施：{x['solution'] or '（未解决）'}"
                     f"｜Root cause：{x['root_cause'] or '待查'}（{x['rc_category'] or '待查'}）｜备件：{x['parts'] or '—'}")
    d = _parse_json(_llm(CONSOLIDATE_PROMPT.format(
        machine=e["machine"], unit=e["unit"], code=f"（故障代码 {e['error_code']}）" if e["error_code"] else "",
        n=len(cases), rcs="、".join(RC_CATEGORIES), cases="\n".join(lines))), kind="dict") or {}
    for lst in (d.get("root_causes") or [], d.get("solutions") or []):
        for it in lst:
            if isinstance(it, dict):
                it["field"], it["refurb"] = int(it.get("field") or 0), int(it.get("refurb") or 0)
                it["count"] = it["field"] + it["refurb"] or int(it.get("count") or 1)
    mix = {}
    for x in cases:
        t = x["exp_type"] or exp_type_of(x["src_doc"])
        mix[t] = mix.get(t, 0) + 1
    meta["exp_mix"] = "·".join(f"{k}{v}" for k, v in sorted(mix.items()))
    meta.update(title=d.get("title") or e["title"],
                root_causes=json.dumps(d.get("root_causes") or [], ensure_ascii=False),
                solutions=json.dumps(d.get("solutions") or [], ensure_ascii=False),
                detours=json.dumps(d.get("detours") or [], ensure_ascii=False),
                follow_up=json.dumps(d.get("process") or [], ensure_ascii=False))
    if e["status"] == "confirmed":
        meta["has_new"] = 1
    c.execute(f"UPDATE ts_entries SET {', '.join(k + '=?' for k in meta)} WHERE id=?", list(meta.values()) + [eid])


def attach_oem(c, eid: int):
    """有故障代码的条目：附上原厂错误码手册 / 要因分析里提到这个码的段落（最多 3 段）。"""
    e = c.execute("SELECT error_code FROM ts_entries WHERE id=?", (eid,)).fetchone()
    code = _code_key(e["error_code"])
    m = re.search(r"[A-Z]{2,5}-[0-9A-F]{1,2}(?:-[0-9A-F]{1,2})?|[0-9A-F]{8}|[0-9A-F]{4}", code)
    if not m:
        return
    from src import acl
    from src.fulltext import search_fts_exact
    tok = acl.set_current_user(acl.SYSTEM_USER)
    try:
        hits = search_fts_exact(m.group(0), limit=12)
    finally:
        acl.reset_current_user(tok)
    refs = [{"doc": h.get("doc_name", ""), "section": h.get("section_title", ""), "text": (h.get("text") or "")[:600]}
            for h in hits if h.get("doc_type") in ("error_code", "manual")][:3]
    c.execute("UPDATE ts_entries SET oem_refs=? WHERE id=?", (json.dumps(refs, ensure_ascii=False), eid))


def refresh_entries(c, eids, log=print):
    for eid in sorted(set(eids)):
        try:
            consolidate(c, eid)
            attach_oem(c, eid)
            c.commit()
        except Exception as ex:
            log(f"  ⚠️ 整理条目 {eid} 失败：{type(ex).__name__}: {ex}")


# ── 4. 主流程：抽取 / 导入 ────────────────────────────────────────────
def run(limit: int = 0, only_doc: str = None, log=print) -> dict:
    c = _conn()
    stat = {"windows": 0, "skipped": 0, "cases": 0, "entries": set()}
    for src_doc, machine, cust, text, section in source_windows(only_doc):
        h = hashlib.md5((src_doc + text).encode()).hexdigest()
        if c.execute("SELECT 1 FROM ts_windows WHERE hash=?", (h,)).fetchone():
            stat["skipped"] += 1
            continue
        if limit and stat["windows"] >= limit:
            break
        stat["windows"] += 1
        try:
            cases = extract_cases(machine, text)
        except Exception as ex:
            log(f"  ⚠️ 抽取失败 {src_doc}：{type(ex).__name__}: {ex}")
            continue
        n_new = 0
        for case in cases:
            eid, changed = add_case(c, machine, case, src_doc, section, cust)
            if changed:
                n_new += 1
                stat["entries"].add(eid)
        c.execute("INSERT OR REPLACE INTO ts_windows(hash, src_doc, n_cases, done_at) VALUES (?,?,?,?)",
                  (h, src_doc, len(cases), _now()))
        c.commit()
        stat["cases"] += n_new
        log(f"  {src_doc[:36]:36s} {section[:28]:28s} → {len(cases)} 条（新 {n_new}）")
    refresh_entries(c, stat["entries"], log)
    backfill_match_keys(c)
    stat["entries"] = len(stat["entries"])
    c.close()
    return stat


def backfill_match_keys(c):
    for r in c.execute("SELECT id, error_code FROM ts_entries WHERE COALESCE(match_key,'')='' AND COALESCE(error_code,'')<>''").fetchall():
        c.execute("UPDATE ts_entries SET match_key=? WHERE id=?", (match_key(r["error_code"]), r["id"]))
    c.commit()


def import_xlsx(path: str, customer: str = "none", log=print) -> dict:
    """导入人工整理好的「报警处理记录」表（列同参考表），算已确认案例。"""
    import openpyxl
    wb = openpyxl.load_workbook(path, data_only=True)
    ws = wb["报警处理记录"] if "报警处理记录" in wb.sheetnames else wb.worksheets[0]
    head = [str(v or "").strip() for v in next(ws.iter_rows(min_row=1, max_row=1, values_only=True))]
    idx = {k: head.index(lbl) for k, lbl in zip(CASE_COLS, CASE_LABELS) if lbl in head}
    c = _conn()
    src = os.path.basename(path)
    eids, n = set(), 0
    for r in ws.iter_rows(min_row=2, values_only=True):
        if not r or not r[0]:
            continue
        case = {}
        for k, i in idx.items():
            v = r[i]
            case[k] = v.strftime("%Y-%m-%d") if hasattr(v, "strftime") else str(v or "").strip()
        if case.get("error_code") in ("-", "—"):
            case["error_code"] = ""
        machine = case.get("machine") or "S207D"
        machine = "S207D" if machine in ("S207", "S207D") else machine
        eid, changed = add_case(c, machine, case, src, case.get("case_no", ""), customer, review="confirmed")
        if changed:
            eids.add(eid)
            n += 1
    c.commit()
    refresh_entries(c, eids, log)
    c.close()
    return {"cases": n, "entries": len(eids)}


def import_oem(json_path: str, log=print) -> dict:
    """导入 ts_parse_oem_guide.py 解析出的原厂 Troubleshooting Guide（origin=oem，直接算已确认）。可重复导入（按 entry_key 覆盖）。"""
    items = json.load(open(json_path, encoding="utf-8"))
    c = _conn()
    n = 0
    for x in items:
        code = x.get("error_code", "")
        key = f"oem|{x['src']}|{x['machine']}|{x['unit']}|{x.get('axis', '')}|{code or _norm_text(x['title'])[:40]}"
        if code and x["src"].startswith("WS RS") and re.match(r"^[0-9A-F]{2}-[0-9A-FX]{2}$", code, re.I):
            code = f"STG-{code.upper()}"     # 原厂表里的 5A-1C = 机台上显示的 STG-5A-1C
        causes = x.get("causes") or []
        rcs = [{"cause": cz["cause"], "category": "", "count": cz.get("rate") or i + 1} for i, cz in enumerate(causes)]
        sols = [{"action": st["check"] + (f"（{st['details']}）" if st.get("details") else ""), "count": 0}
                for cz in causes for st in cz.get("steps", []) if st.get("check")]
        title = x["title"] + (f"｜{x['axis']} 轴" if x.get("axis") else "")
        c.execute("""INSERT INTO ts_entries(entry_key, machine, unit, error_code, title, root_causes, solutions,
                         origin, oem_detail, src, match_key, status, occ_count, updated_at)
                     VALUES (?,?,?,?,?,?,?, 'oem', ?,?,?, 'confirmed', 0, ?)
                     ON CONFLICT(entry_key) DO UPDATE SET title=excluded.title, root_causes=excluded.root_causes,
                         solutions=excluded.solutions, oem_detail=excluded.oem_detail, updated_at=excluded.updated_at""",
                  (key, x["machine"], x["unit"], code, title, json.dumps(rcs, ensure_ascii=False),
                   json.dumps(sols, ensure_ascii=False), json.dumps(causes, ensure_ascii=False), x["src"],
                   match_key(code), _now()))
        n += 1
    c.commit()
    backfill_match_keys(c)      # 补上加字段之前建的现场条目的匹配键
    c.close()
    log(f"导入原厂条目 {n} 个")
    return {"entries": n}


def related(eid: int, user=None, admin=False) -> list:
    """同机型、同错误码的另一来源条目（现场 ↔ 原厂）。"""
    c = _conn()
    e = c.execute("SELECT * FROM ts_entries WHERE id=?", (eid,)).fetchone()
    if not e or not e["match_key"]:
        c.close()
        return []
    rows = c.execute("SELECT * FROM ts_entries WHERE machine=? AND match_key=? AND id<>? AND origin<>?",
                     (e["machine"], e["match_key"], eid, e["origin"] or "field")).fetchall()
    out = [dict(r) for r in rows if _visible(user, r, admin, c)]
    c.close()
    return out


def alarm_view(e: dict, rel: list, md: bool = True) -> str:
    """同一个报警的统一视图：Root cause（全部列出，标来源）+ 处理动作（原厂推荐 / 现场有效 / 翻新有效）+ 少走弯路。
    e 为条目（现场或原厂），rel 为同机型同错误码的另一来源条目。"""
    allx = [e] + list(rel)
    oem = [x for x in allx if (x.get("origin") or "field") == "oem"]
    fld = [x for x in allx if (x.get("origin") or "field") != "oem"]
    b = (lambda t: f"**{t}**") if md else (lambda t: t)
    out = []
    occ = sum(int(x.get("occ_count") or 0) for x in fld)
    mix = "、".join(x.get("exp_mix") or "" for x in fld if x.get("exp_mix"))
    out.append(f"{e['machine']}｜{e['unit']}｜故障代码：{e['error_code'] or '—'}　"
               + (f"现场+翻新共记录 {occ} 次（{mix}）" if fld else "现场 / 翻新还没有记录")
               + ("　＋原厂 Guide" if oem else ""))
    # Root cause
    rc_lines = []
    for o in oem:
        for cz in json.loads(o.get("oem_detail") or "[]"):
            rc_lines.append(f"【原厂·频度第{cz.get('rate')}】{cz.get('cause')}")
    for f in fld:
        for r in json.loads(f.get("root_causes") or "[]"):
            fc, rf = int(r.get("field") or 0), int(r.get("refurb") or 0)
            tag = "".join(t for t in (f"【现场 {fc} 次】" if fc else "", f"【翻新 {rf} 次】" if rf else "")) \
                or f"【记录 {r.get('count', 1)} 次】"
            rc_lines.append(f"{tag}{r.get('cause')}" + (f"（{r.get('category')}）" if r.get("category") else ""))
    out.append("")
    out.append(b(f"Root cause（共 {len(rc_lines)} 个）"))
    out += [f"{i}. {t}" for i, t in enumerate(rc_lines, 1)] or ["- 待查"]
    # 处理动作
    out.append("")
    out.append(b("处理动作"))
    if oem:
        out.append(b("▸ 原厂推荐（Nikon Troubleshooting Guide）"))
        for o in oem:
            for cz in json.loads(o.get("oem_detail") or "[]"):
                for st in cz.get("steps", []):
                    out.append(f"- [针对原因{cz.get('rate')}] {st.get('check') or '—'}"
                               + (f"：{st['details'][:200]}" if st.get("details") else ""))
            if not json.loads(o.get("oem_detail") or "[]"):
                out.append("- Guide 只列了错误名称，没有写对策")
    for kind, key in (("现场验证有效", "field"), ("翻新验证有效", "refurb")):
        acts = [a for f in fld for a in json.loads(f.get("solutions") or "[]") if int(a.get(key) or 0) > 0]
        if acts:
            out.append(b(f"▸ {kind}"))
            out += [f"- {a.get('action')}（解决 {a.get(key)} 次" + (f"；针对：{a['for_cause']}" if a.get("for_cause") else "") + "）"
                    for a in sorted(acts, key=lambda a: -int(a.get(key) or 0))]
    legacy = [a for f in fld for a in json.loads(f.get("solutions") or "[]")
              if not int(a.get("field") or 0) and not int(a.get("refurb") or 0) and a.get("action")]
    if legacy:
        out.append(b("▸ 有效措施（来源未区分）"))
        out += [f"- {a.get('action')}（解决 {a.get('count', 1)} 次）" for a in legacy]
    dets = [d for f in fld for d in json.loads(f.get("detours") or "[]")]
    if dets:
        out.append(b("▸ 少走弯路（试过但没用）"))
        out += [f"- {d}" for d in dets]
    return "\n".join(out)


# ── 5. 查询 / 可见性 / 编辑（网页和智能体用）────────────────────────────
def _visible(user, e, admin: bool, conn=None) -> bool:
    from src import acl
    if admin or not acl.security_enabled():
        return True
    if e["status"] != "confirmed":
        return False
    if (e["origin"] or "field") == "oem":
        return acl.can_see(user, "oem_manual", e["machine"])
    if not acl.can_see(user, "fault_history", e["machine"]):
        return False
    custs = [x for x in (e["customers"] or "").split(",") if x]
    if custs:
        ucust = user.grants[2] if user is not None else set()
        if not (set(custs) & set(ucust)):
            c = conn or _conn()
            n_plain = c.execute("SELECT count(*) FROM ts_cases WHERE entry_id=? AND customer IN ('', 'none')",
                                (e["id"],)).fetchone()[0]
            if conn is None:
                c.close()
            if n_plain == 0:
                return False      # 全部案例来自没权限的客户资料
    return True


def search(user=None, admin=False, machine="", unit="", status="", keyword="", want_models=None, limit=5000,
           origin="") -> list:
    c = _conn()
    sql, params = "SELECT * FROM ts_entries WHERE 1=1", []
    if origin:
        sql += " AND COALESCE(origin,'field')=?"; params.append(origin)
    if machine:
        sql += " AND machine=?"; params.append(machine)
    if unit:
        sql += " AND unit=?"; params.append(unit)
    if status == "has_new":
        sql += " AND has_new=1"
    elif status:
        sql += " AND status=?"; params.append(status)
    rows = c.execute(sql + " ORDER BY (COALESCE(origin,'field')='oem'), occ_count DESC, last_date DESC, id", params).fetchall()
    from src.error_code import from_screen
    scr = from_screen(keyword or "")
    if scr:                       # 搜索框里贴了报警界面原文 → 按换算出的 STG-10-17 搜
        keyword = scr[0]
    kw = [k.lower() for k in re.split(r"\s+", (keyword or "").strip()) if k]
    out = []
    for e in rows:
        if want_models and e["machine"] not in want_models:
            continue
        if not _visible(user, e, admin, c):
            continue
        if kw:
            hay = " ".join(str(e[k] or "") for k in ("error_code", "match_key", "title", "root_causes", "solutions", "unit")).lower()
            if not all(k in hay for k in kw):
                continue
        out.append(dict(e))
        if len(out) >= limit:
            break
    c.close()
    return out


def entry_detail(eid: int, user=None, admin=False) -> dict:
    c = _conn()
    e = c.execute("SELECT * FROM ts_entries WHERE id=?", (eid,)).fetchone()
    if not e or not _visible(user, e, admin, c):
        c.close()
        return {}
    cases = [dict(x) for x in c.execute("SELECT * FROM ts_cases WHERE entry_id=? ORDER BY reg_date, id", (eid,))]
    c.close()
    if not admin:
        ucust = user.grants[2] if user is not None else set()
        cases = [x for x in cases if x["customer"] in ("", "none") or x["customer"] in ucust]
    return {"entry": dict(e), "cases": cases}


def save_entry(admin: str, eid: int, **fields) -> str:
    allowed = {"title", "unit", "machine", "error_code", "root_causes", "solutions", "detours", "follow_up",
               "notes", "status"}
    sets = {k: v for k, v in fields.items() if k in allowed and v is not None}
    if not sets:
        return "⚠️ 没有要保存的内容。"
    sets.update(edited=1, has_new=0, edited_by=admin, updated_at=_now())
    c = _conn()
    c.execute(f"UPDATE ts_entries SET {', '.join(k + '=?' for k in sets)} WHERE id=?", list(sets.values()) + [eid])
    if sets.get("status") == "confirmed":
        c.execute("UPDATE ts_cases SET review='confirmed' WHERE entry_id=?", (eid,))
    c.commit()
    c.close()
    try:
        from src import audit
        audit.log(admin, f"改 Troubleshooting 表 #{eid}：{','.join(k for k in fields if fields[k] is not None)}"[:200], ok=True)
    except Exception:
        pass
    return f"✅ 已保存条目 #{eid}" + ("，已确认（有权限的账号可见、智能体会引用）" if sets.get("status") == "confirmed" else "")


def lookup_for_agent(query: str, user, want_models=None, k: int = 4) -> str:
    """智能体用：按故障代码 / 关键词找已确认条目，输出现场验证过的根因、解决措施、弯路，以及原厂 Guide。
    问题里有错误码时只认码对得上的（同码 > 同报警名称 > 同主码参考），不再按关键词乱配。"""
    from src.error_code import annotate, parse, lookup as code_lookup
    q = annotate((query or "").strip())
    codes = re.findall(r"[A-Z]{2,5}-[0-9A-F]{1,2}(?:-[0-9A-F]{1,2})?|#?[0-9A-F]{2}H?-[0-9A-F]{2}|[0-9A-F]{8}|[0-9A-F]{4}", q.upper())
    mkeys = {match_key(cd) for cd in codes} - {""}
    mains = {mk.split("-")[0] for mk in mkeys if re.match(r"^[0-9A-F]{2}-", mk)}
    # 报警名称：STG 码换算后从原文 / 手册取主码名称（现场记录常只抄了报警文字、没记数字）
    names = []
    for ctrl, main, det in parse(q):
        txt = code_lookup(f"{ctrl}-{main}-{det}" if det else f"{ctrl}-{main}") or ""
        m = re.search(rf"#{main}H[^）)]*[）)]：\s*(.+?)(?:\(|（|\n|$)", txt)   # 必须是这个主码自己的名称
        words = [w for w in re.findall(r"[a-z0-9]+", (m.group(1) if m else "").lower()) if len(w) >= 3]
        if len(words) >= 2:
            names.append(words)       # 报警名称里 3 个字母以上的词要全部出现（容忍「No return FP (WT) response」这种插字）
    hits = []
    for e in search(user, admin=False, status="confirmed", want_models=want_models):
        score, why = 0, ""
        ec = _code_key(e["error_code"])
        hay = " ".join(str(e[f] or "") for f in ("title", "root_causes", "solutions", "error_code", "unit")).lower()
        hay_words = set(re.findall(r"[a-z0-9]+", hay))
        if e["match_key"] and e["match_key"] in mkeys:
            score, why = 12, ""
        elif ec and any(cd in ec or ec in cd for cd in codes if len(cd) >= 4):
            score, why = 10, ""
        elif names and (e["origin"] or "field") != "oem" and not re.match(r"^[0-9A-F]{2}-", e["match_key"] or "") \
                and any(all(w in hay_words for w in ws) for ws in names):
            # 按报警名称只配「现场只抄了报警文字、没记数字代码」的记录；原厂 / 有数字码的条目各有准确代码
            score, why = 8, "（同报警名称，现场当时没记数字代码）"
        elif mains and e["match_key"][:2] in mains and re.match(r"^[0-9A-F]{2}-", e["match_key"] or ""):
            score, why = 3, f"（同主码 {e['match_key'][:2]}，仅供参考）"
        elif not codes:
            score = sum(1 for w in re.split(r"[\s,，、]+", q.lower()) if len(w) >= 2 and w in hay)
        if score:
            e = dict(e); e["_why"] = why
            hits.append((score, e))
    hits.sort(key=lambda x: (-x[0], -(x[1]["occ_count"] or 0)))
    if any(sc >= 8 for sc, _ in hits):                 # 有同码 / 同报警名称的，就不列同主码参考
        hits = [(sc, e) for sc, e in hits if sc >= 8]
    shown = sorted(k for k in mkeys if "-" in k)
    if not hits:
        return (f"Troubleshooting 表里没有 {'、'.join(shown)} 的已确认条目。") if shown \
            else "Troubleshooting 表里没有已确认的相关条目。"
    out, seen = [], set()
    for _, e in hits:
        gk = (e["machine"], e["match_key"]) if e["match_key"] else ("id", e["id"])
        if gk in seen:
            continue
        seen.add(gk)
        rel = related(e["id"], user, False)
        head = "Troubleshooting 表" if (e["origin"] or "field") != "oem" else "原厂 Troubleshooting Guide"
        out.append(f"【{head}】{e['title']}{e.get('_why', '')}\n" + alarm_view(e, rel, md=False))
        if len(out) >= k:
            break
    return "\n\n".join(out)
