"""机型标签：入库时先读内容判断，判断不出来再按 WPS 路径/文件名推断；检索时按机型加权（只降权、不过滤）。

标签用机型"族"：
  S206D 族：S205C / S206D / S305B / S306C / S306D
  S207D 族：S207D / S307E
  SF100 族：SF100 ~ SF140 / SF200
  S308F 族：S308F（X6 平台：AR-X6 / EX-X6）——2026-10-15 加
通用文档可带多个标签，逗号分隔，如 "S206D,S207D"。
环境变量：
  DEFAULT_MACHINE_MODEL  无法推断时的默认标签（S207D）
  PRIMARY_MODELS         问题没点名机型时的优先顺序（S207D,S206D），依次乘 1.0 / 0.8 …
  OTHER_MODEL_FACTOR     非优先机型 / 与问题点名机型不符时的系数（0.5）
"""
import os
import re

# 文件名、标签用的严格规则
STRICT = [
    ("S206D", re.compile(r"S20[56][CD]|S30[56][BCD]", re.I)),
    ("S207D", re.compile(r"S207D|S307E", re.I)),
    ("SF100", re.compile(r"SF1[0-4]0|SF200", re.I)),
    ("S308F", re.compile(r"S308F|(?<![A-Z0-9])(?:AR|EX)-?X6(?![0-9])|(?<![A-Z0-9])X6(?![0-9])", re.I)),
]
# 用户提问用的宽松规则（"207""206 的""S307" 这类说法）
LOOSE = [
    ("S206D", re.compile(r"(?<![0-9A-Z])(?:NSR-?)?S?20[56][CD]?(?![0-9])|(?<![0-9A-Z])S?30[56][BCD]?(?![0-9])", re.I)),
    ("S207D", re.compile(r"(?<![0-9A-Z])(?:NSR-?)?S?207D?(?![0-9])|(?<![0-9A-Z])S?307E?(?![0-9])", re.I)),
    ("SF100", re.compile(r"SF\s?1[0-4]0|SF\s?200", re.I)),
    ("S308F", re.compile(r"(?<![0-9A-Z])(?:NSR-?)?S?308F?(?![0-9])|(?<![A-Z0-9])(?:AR|EX)?-?X6(?![0-9])", re.I)),
]
DIR_FAMILY = {"206": "S206D", "207": "S207D", "307": "S207D", "SF100": "SF100", "308": "S308F"}
ORDER = ["COMMON", "S207D", "S206D", "SF100", "S308F"]          # 标签排列顺序


def _common_dirs():
    """跟机型无关的通用资料所在目录（关键字，逗号分隔）。带 COMMON 标签的文档任何问题都不降权。"""
    return [d.strip() for d in os.getenv("COMMON_DIRS", "光刻原理").split(",") if d.strip()]


def _sort_tags(tags):
    tags = [t for t in dict.fromkeys(tags) if t]
    return sorted(tags, key=lambda t: ORDER.index(t) if t in ORDER else len(ORDER))


def _default():
    return os.getenv("DEFAULT_MACHINE_MODEL", "S207D")


def _primary():
    return [m.strip() for m in os.getenv("PRIMARY_MODELS", "S207D,S206D").split(",") if m.strip()]


def _families(text, rules=STRICT):
    return [tag for tag, pat in rules if pat.search(text or "")]


def infer_model(rel_path: str = "", name: str = "") -> str:
    """WPS 路径上最近的机型目录优先；其次文件名；都没有给默认值。"""
    if any(d in (rel_path or "") for d in _common_dirs()):
        return "COMMON"
    parts = [p for p in re.split(r"[\\/]", rel_path or "") if p]
    for p in reversed(parts[:-1] if parts else []):
        if p in DIR_FAMILY:
            return DIR_FAMILY[p]
    found = _sort_tags(_families(name))
    return ",".join(found) if found else _default()


def normalize(tags: str) -> list:
    """把库里的标签（含旧写法 NSR-S207D）归一成族名列表。"""
    out = []
    for raw in re.split(r"[,，;/ ]+", tags or ""):
        if not raw:
            continue
        fams = _families(raw)
        out.extend(fams if fams else [raw.upper()])
    return list(dict.fromkeys(out))


def models_in_query(query: str) -> list:
    return _families(query, LOOSE)


def model_factor(tags: str, query: str = "") -> float:
    fams = normalize(tags)
    if not fams or "COMMON" in fams:
        return 1.0                      # 没标签的、通用资料不动
    other = float(os.getenv("OTHER_MODEL_FACTOR", "0.5"))
    want = models_in_query(query)
    if want:
        return 1.0 if any(f in want for f in fams) else other
    prim = _primary()
    best = other
    for f in fams:
        if f in prim:
            best = max(best, 1.0 - 0.2 * prim.index(f))
    return best


# ── 按内容判断机型（入库时优先用；判断不出来再退回路径/文件名）──────────────
# 正文里数机型名要加边界，免得把 4S306C… 这类部件号算进去
CONTENT = [
    ("S206D", re.compile(r"(?<![0-9A-Z])(?:NSR-?)?(?:S20[56][CD]|S30[56][BCD])(?![0-9])", re.I)),
    ("S207D", re.compile(r"(?<![0-9A-Z])(?:NSR-?)?(?:S207D|S307E)(?![0-9])", re.I)),
    ("SF100", re.compile(r"(?<![0-9A-Z])(?:NSR-?)?(?:SF1[0-4]0|SF200)(?![0-9])", re.I)),
    ("S308F", re.compile(r"(?<![0-9A-Z])(?:NSR-?)?S308F(?![0-9])|(?<![A-Z0-9])(?:AR|EX)-?X6(?![0-9])", re.I)),
]
# Nikon 文档编号，如 S307E-RF-ELCIRC-R10、S206D-SM-…：编号里的机型就是这份资料的机型
DOC_NO = re.compile(r"(?<![0-9A-Z])((?:S20[567][CD]|S30[567][BCDE]|S308F|SF1[0-4]0|SF200))-[A-Z]{2,}-", re.I)
COVER_CHARS = 3000          # Word/Excel/文本：开头这么多字算"封面"
BODY_CHARS = 200_000        # 正文抽样上限
BODY_PAGES = 40             # PDF 正文最多均匀抽这么多页
MIN_SCORE = 3               # 证据太少（扫描件、几乎没字）就不下结论


def _sample_text(path) -> tuple:
    """返回 (封面文字, 正文抽样文字)。读不了的格式返回两个空串。"""
    from pathlib import Path
    path = Path(path)
    suf = path.suffix.lower()
    try:
        if suf == ".pdf":
            import fitz
            with fitz.open(str(path)) as doc:
                n = doc.page_count
                cover = "\n".join(doc[i].get_text() for i in range(min(2, n)))
                rest = list(range(2, n))
                if len(rest) > BODY_PAGES:
                    step = len(rest) / BODY_PAGES
                    rest = [rest[int(i * step)] for i in range(BODY_PAGES)]
                body = "\n".join(doc[i].get_text() for i in rest)
            return cover, body[:BODY_CHARS]
        if suf == ".docx":
            import docx
            d = docx.Document(str(path))
            parts = [d.core_properties.title or "", d.core_properties.subject or ""]
            for sec in d.sections:                       # 页眉页脚里常有机型和文档编号
                for hf in (sec.header, sec.footer):
                    parts.extend(pp.text for pp in hf.paragraphs)
            parts.extend(pp.text for pp in d.paragraphs)
            for t in d.tables:
                for row in t.rows:
                    parts.extend(c.text for c in row.cells)
                if sum(map(len, parts)) > BODY_CHARS:
                    break
            text = "\n".join(x for x in parts if x)
            return text[:COVER_CHARS], text[COVER_CHARS:BODY_CHARS]
        if suf == ".xlsx":
            import openpyxl
            wb = openpyxl.load_workbook(str(path), read_only=True, data_only=True)
            parts, size = list(wb.sheetnames), 0
            for ws in wb.worksheets:
                for row in ws.iter_rows(values_only=True):
                    line = " ".join(str(v) for v in row if v is not None)
                    if line:
                        parts.append(line); size += len(line)
                    if size > BODY_CHARS:
                        break
                if size > BODY_CHARS:
                    break
            wb.close()
            text = "\n".join(parts)
            return text[:COVER_CHARS], text[COVER_CHARS:BODY_CHARS]
        if suf in (".md", ".txt", ".csv"):
            text = path.read_text(errors="ignore")[:BODY_CHARS]
            return text[:COVER_CHARS], text[COVER_CHARS:]
    except Exception as e:
        print(f"    （读内容判断机型失败，退回路径规则：{type(e).__name__}: {e}）")
    return "", ""


def infer_model_from_content(path) -> tuple:
    """读文件内容判断机型。返回 (标签 或 "", 依据说明)。

    依据从强到弱：
      1. 封面（PDF 前 2 页 / 文档开头）写的机型、Nikon 文档编号里的机型 —— 有就只认这些；
      2. 否则数正文抽样里各机型出现的次数，取最多的那个，
         次数达到最多者 30% 以上的也一起标上（真正的多机型通用资料）。
    证据不足返回 ("", 原因)，由调用方退回路径/文件名规则。
    """
    cover, body = _sample_text(path)
    if not (cover or body).strip():
        return "", "抽不出文字"
    strong = {}
    for m in DOC_NO.finditer(cover + "\n" + body):
        for fam in _families(m.group(1)):
            strong[fam] = strong.get(fam, 0) + 1
    for fam, pat in CONTENT:
        n = len(pat.findall(cover))
        if n:
            strong[fam] = strong.get(fam, 0) + n
    if strong:
        tags = _sort_tags(strong)
        return ",".join(tags), "封面/文档编号：" + "，".join(f"{k}×{v}" for k, v in strong.items())
    counts = {fam: len(pat.findall(body)) for fam, pat in CONTENT}
    counts = {k: v for k, v in counts.items() if v}
    if not counts or max(counts.values()) < MIN_SCORE:
        return "", "正文机型字样太少" + (f"（{counts}）" if counts else "")
    top = max(counts.values())
    tags = _sort_tags([k for k, v in counts.items() if v >= max(MIN_SCORE, 0.3 * top)])
    return ",".join(tags), "正文次数：" + "，".join(f"{k}×{v}" for k, v in counts.items())


def infer_model_for_file(path, rel_paths=None, name: str = "") -> tuple:
    """入库用：通用目录 → COMMON；否则先按内容判断，判断不出来再按路径/文件名。返回 (标签, 依据)。"""
    from pathlib import Path
    name = name or Path(path).name
    if any(d in (rp or "") for rp in (rel_paths or []) for d in _common_dirs()):
        return "COMMON", "通用资料目录"
    tag, why = infer_model_from_content(path)
    if tag:
        return tag, why
    return infer_model_multi(rel_paths, name), f"按路径/文件名（{why}）"


def infer_model_multi(rel_paths, name: str = "") -> str:
    """同一文件出现在多个机型目录（如 206、307、SF100 各放一份）时，合并各路径的标签。"""
    tags = []
    for rp in (rel_paths or [""]):
        tags.extend(infer_model(rp, name).split(","))
    tags = _sort_tags(tags)
    return ",".join(tags) if tags else _default()


def mirror_paths(mirror, name: str) -> list:
    """镜像目录里所有同名文件的相对路径。"""
    import os
    root = str(mirror)
    out = []
    for d, _, files in os.walk(root):
        if name in files:
            out.append(os.path.relpath(os.path.join(d, name), root))
    return out
